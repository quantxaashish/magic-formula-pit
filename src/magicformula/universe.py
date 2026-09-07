"""NSE/BSE universe construction and AMFI cap-bucket tagging (SPEC.md
section 1).

Four stages:
  1. Fetch + parse NSE's equity list and BSE's active equity scrip list
     (each a real, official public archive - safe to use freely per
     SPEC.md section 2.5, unlike screener.in).
  2. Merge on ISIN (a company can be listed on both exchanges; this
     project treats it as one entity, preferring the NSE symbol/listing
     date when both exist, since NSE has deeper liquidity - SPEC.md
     section 1).
  3. Tag each merged entity with its AMFI cap bucket (large/mid/small),
     parsed from AMFI's own published categorization file.
  4. Apply the exclusions below - including a name-heuristic pre-filter
     for Financial Services/Utilities (SPEC.md section 5) - so that
     fundamentals.py's fetch layer never wastes a scrape (and retries) on
     a company that would be excluded from ranking anyway. This ordering
     matters: banks/NBFCs/insurers use a completely different screener.in
     page template (Interest Earned/Expended, no "Operating Profit" row),
     so fetching them first and filtering after means paying the full
     retry cost for every one of them, then discarding the result.

What this module does NOT yet do (documented, not silently assumed):
  - Full-fidelity exclusion of ETFs/REITs/InvITs/preference shares/
    partly-paid shares and SME-platform listings. NSE's mainboard
    EQUITY_L.csv and BSE's segment=Equity filter already exclude some of
    these by construction (ETFs and NSE Emerge/SME listings aren't part
    of either endpoint at all), but this hasn't been verified line-by-line
    against NSE/BSE's own REIT/InvIT/preference-share reference lists.
    exclude_non_standard_instruments() below is a name-pattern heuristic,
    not an authoritative filter - flag anything it catches for a human
    second look rather than trusting it blindly.
  - Full-fidelity Financial Services/Utilities exclusion.
    exclude_likely_financials_and_utilities() is also a name-pattern
    heuristic (a pre-filter for fetch cost, see stage 4 above) - a company
    that slips past it still needs a real sector-based exclusion once
    fetched, which is fundamentals.py's/a later pipeline stage's job, not
    this module's.
  - Surveillance/trade-to-trade (ASM/GSM) watchlist flagging. NSE's
    "SERIES" column (BE/BZ here mean trade-to-trade settlement, not
    necessarily surveillance) is captured as metadata but not used to
    exclude anything - a real surveillance flag needs NSE/BSE's dedicated
    ASM/GSM lists, not inferred from settlement series. THIS IS A HARD
    BLOCKER FOR THE SMALL-CAP BUCKET SPECIFICALLY before any real position
    sizing: small/micro caps are where ASM/GSM surveillance flags and
    thin, manipulable liquidity concentrate, and this pipeline currently
    has no way to detect either. See README's Known Limitations.
  - Insolvency (NCLT) status - no data source wired up yet. Same
    small-cap-specific blocker as above: distressed small caps are exactly
    where this would matter most and exactly where it's least likely to
    be caught by any of the other filters here.

Compliance (SPEC.md section 2.5): NSE/BSE equity lists and AMFI's
categorization file are official public archives, safe to use freely for
this purpose - no rate limiting or caching discipline is required here the
way it is for screener.in, though caching the versioned snapshots to disk
(as SPEC.md section 1 asks) still avoids needless re-fetching.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Literal

import openpyxl
import requests

logger = logging.getLogger(__name__)

DEFAULT_USER_AGENT = "MagicFormulaIndia-Research/0.1 (personal, non-commercial research tool)"

NSE_EQUITY_LIST_URL = "https://nsearchives.nseindia.com/content/equities/EQUITY_L.csv"
BSE_ACTIVE_EQUITY_API_URL = (
    "https://api.bseindia.com/BseIndiaAPI/api/ListofScripData/w"
    "?Group=&Scripcode=&industry=&segment=Equity&status=Active"
)
# The categorization file itself is dated by its 6-month data snapshot
# (e.g. "31Dec2025" = average market cap over Jul-Dec 2025) but takes
# effect for the *following* 6 months (Jan-Jun 2026 in that example).
# AMFI republishes this page with a new file roughly every January and
# July - there's no fixed, predictable URL for "the current edition", so
# this always needs a human (or a scheduled job) to check the page and
# pass the current file's URL in explicitly.
AMFI_CATEGORIZATION_PAGE_URL = "https://www.amfiindia.com/otherdata/categorisation-of-stocks"

Exchange = Literal["NSE", "BSE"]
CapBucket = Literal["large", "mid", "small"]

# Heuristic only - see module docstring. Matches on whole-word patterns in
# the company name for instrument types that shouldn't be in the Magic
# Formula universe at all (SPEC.md section 1), not on any authoritative
# NSE/BSE instrument-type field (neither source exposes one in the
# endpoints this module uses). Word-boundary regex, not substring
# containment: a plain "TRUST" substring check flags names like "Trustwave
# Securities" that merely contain the letters, and an earlier "-RE" marker
# (dropped entirely - see decision 0002) flagged names ending in "-Regular"
# or "-Reinvestment" (mutual fund unit variants) with zero real REIT/InvIT
# catches to show for it: real REITs/InvITs (Embassy Office Parks,
# Mindspace, Brookfield India REIT, IRB InvIT, IndiGrid, PowerGrid InvIT,
# IndInfravit, Bharat Highways InvIT - checked directly against the live
# universe) don't appear in NSE's EQUITY_L.csv or BSE's segment=Equity feed
# at all; they trade under a different instrument type on both exchanges,
# so this data source excludes them structurally before this heuristic
# ever runs. What this heuristic actually catches in practice is a
# handful of "Investment Trust"/"Capital Trust"-style financial companies
# that use "Trust" in their name without being REITs - a defensible
# exclusion in spirit (they're investment vehicles, not operating
# companies) but mislabeled if called a REIT/InvIT catch, hence the
# exclusion_reason wording below.
NON_STANDARD_NAME_PATTERN = re.compile(r"\b(REIT|INVIT|INVITS|TRUST)\b", re.IGNORECASE)

# Heuristic only, same caveat as NON_STANDARD_NAME_MARKERS - this exists to
# pre-filter financials/utilities (SPEC.md section 5) *before* they ever
# reach fundamentals.py's fetch layer, since their screener.in page uses a
# completely different template (Interest Earned/Expended for banks, no
# "Operating Profit" row at all) that fundamentals.py's parser doesn't
# handle and isn't meant to - there's no reason to spend a scrape (and
# retries) on a company that section 5 excludes from ranking regardless.
# This is a pre-filter for cost, not the authoritative section 5
# exclusion - a company that slips past this heuristic and gets fetched
# anyway should still be excluded properly once its real sector is known.
# "CAPITAL", "CREDIT" and "PAYMENT" were added after a 236-company pilot
# run showed real misses: Aditya Birla Capital (NBFC), SBI Cards and
# Payment Services, and REC Limited all slipped past the narrower marker
# set (see docs/decisions/0003). "REC Limited" (a well-known NBFC, formerly
# Rural Electrification Corporation) still won't be caught by any name
# heuristic - its short registered name carries no financial-sounding word
# at all. That's an accepted gap, not something to special-case by
# hardcoding one company's name.
FINANCIAL_NAME_MARKERS = (
    "BANK", "FINANCE", "FINANCIAL", "NBFC", "INSURANCE", "ASSET MANAGEMENT",
    "HOUSING FIN", "MUTUAL FUND", "CHIT FUND", "MICROFINANCE", "LEASING",
    "CAPITAL", "CREDIT", "PAYMENT",
)
UTILITY_NAME_MARKERS = (
    "POWER", "ELECTRIC", "TRANSMISSION", "DISCOM", "UTILITIES", "GRID", " GAS",
)


@dataclass
class NseListing:
    symbol: str
    name: str
    series: str
    date_of_listing: date
    isin: str
    face_value: float


@dataclass
class BseListing:
    scrip_code: str
    symbol: str | None  # BSE's own ticker (scrip_id); can be missing
    name: str
    isin: str
    group: str


@dataclass
class UniverseEntry:
    isin: str
    name: str
    primary_exchange: Exchange
    primary_symbol: str
    nse_symbol: str | None
    bse_symbol: str | None
    bse_scrip_code: str | None
    date_of_listing: date | None  # from NSE only; BSE listing dates aren't in the feed used here
    nse_series: str | None = None
    cap_bucket: CapBucket | None = None
    cap_rank: int | None = None
    excluded: bool = False
    exclusion_reason: str | None = None


# --- NSE ------------------------------------------------------------------


def parse_nse_equity_csv(text: str) -> list[NseListing]:
    """Pure parser for NSE's EQUITY_L.csv - no network I/O."""
    reader = csv.DictReader(io.StringIO(text))
    listings = []
    for row in reader:
        row = {k.strip(): v.strip() for k, v in row.items()}
        listings.append(
            NseListing(
                symbol=row["SYMBOL"],
                name=row["NAME OF COMPANY"],
                series=row["SERIES"],
                date_of_listing=datetime.strptime(row["DATE OF LISTING"], "%d-%b-%Y").date(),
                isin=row["ISIN NUMBER"],
                face_value=float(row["FACE VALUE"]),
            )
        )
    return listings


def fetch_nse_equity_list(
    session: requests.Session | None = None, timeout: float = 20.0
) -> list[NseListing]:
    session = session or requests.Session()
    response = session.get(
        NSE_EQUITY_LIST_URL, headers={"User-Agent": DEFAULT_USER_AGENT}, timeout=timeout
    )
    response.raise_for_status()
    return parse_nse_equity_csv(response.text)


# --- BSE --------------------------------------------------------------------


def parse_bse_equity_json(data: list[dict]) -> list[BseListing]:
    """Pure parser for BSE's ListofScripData API response - no network I/O."""
    return [
        BseListing(
            scrip_code=row["SCRIP_CD"],
            symbol=row.get("scrip_id") or None,
            name=row["Scrip_Name"],
            isin=row["ISIN_NUMBER"],
            group=row["GROUP"],
        )
        for row in data
    ]


def fetch_bse_active_equity_list(
    session: requests.Session | None = None, timeout: float = 20.0
) -> list[BseListing]:
    session = session or requests.Session()
    response = session.get(
        BSE_ACTIVE_EQUITY_API_URL,
        headers={
            "User-Agent": DEFAULT_USER_AGENT,
            "Referer": "https://www.bseindia.com/",
            "Accept": "application/json, text/plain, */*",
        },
        timeout=timeout,
    )
    response.raise_for_status()
    return parse_bse_equity_json(response.json())


# --- Merge (ISIN-keyed, prefer NSE) -----------------------------------------


def merge_universe(nse_listings: list[NseListing], bse_listings: list[BseListing]) -> list[UniverseEntry]:
    """Merge NSE and BSE listings into one entity per ISIN, preferring the
    NSE symbol and listing date when a company is on both exchanges
    (SPEC.md section 1: NSE has deeper liquidity).
    """
    bse_by_isin: dict[str, BseListing] = {}
    for bse in bse_listings:
        if bse.isin in bse_by_isin:
            logger.warning(
                "BSE: duplicate ISIN %s (%s and %s) - keeping the first seen",
                bse.isin, bse_by_isin[bse.isin].name, bse.name,
            )
            continue
        bse_by_isin[bse.isin] = bse

    entries: list[UniverseEntry] = []
    seen_isins: set[str] = set()

    for nse in nse_listings:
        if nse.isin in seen_isins:
            logger.warning("NSE: duplicate ISIN %s (%s) - keeping the first seen", nse.isin, nse.name)
            continue
        seen_isins.add(nse.isin)
        bse = bse_by_isin.get(nse.isin)
        entries.append(
            UniverseEntry(
                isin=nse.isin,
                name=nse.name,
                primary_exchange="NSE",
                primary_symbol=nse.symbol,
                nse_symbol=nse.symbol,
                bse_symbol=bse.symbol if bse else None,
                bse_scrip_code=bse.scrip_code if bse else None,
                date_of_listing=nse.date_of_listing,
                nse_series=nse.series,
            )
        )

    for bse in bse_listings:
        if bse.isin in seen_isins:
            continue
        seen_isins.add(bse.isin)
        entries.append(
            UniverseEntry(
                isin=bse.isin,
                name=bse.name,
                primary_exchange="BSE",
                primary_symbol=bse.symbol or bse.scrip_code,
                nse_symbol=None,
                bse_symbol=bse.symbol,
                bse_scrip_code=bse.scrip_code,
                date_of_listing=None,
            )
        )

    return entries


# --- AMFI cap-bucket categorization ------------------------------------------


def parse_amfi_categorization(xlsx_bytes: bytes) -> dict[str, tuple[int, CapBucket]]:
    """Pure parser for AMFI's categorization Excel file - no network I/O.
    Returns {isin: (rank, cap_bucket)}. Rank 1-100 = large, 101-250 = mid,
    251+ = small, per SEBI circular SEBI/HO/IMD/DF3/CIR/P/2017/114 (the
    file's own "Categorization" column already applies this, so this
    parser trusts AMFI's column rather than re-deriving it from rank -
    the rank is kept only for reference/debugging).
    """
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    category_map: dict[str, CapBucket] = {
        "Large Cap": "large", "Mid Cap": "mid", "Small Cap": "small",
    }

    rows = sheet.iter_rows(min_row=1, values_only=True)
    header = next(rows)
    while header and header[1] != "Company name":
        header = next(rows)  # skip title row(s) until the real header

    col = {name: i for i, name in enumerate(header) if name}
    result: dict[str, tuple[int, CapBucket]] = {}
    for row in rows:
        if not row or row[col["ISIN"]] is None:
            continue
        isin = str(row[col["ISIN"]]).strip()
        rank = row[col["Sr. No."]]
        category_label = row[col["Categorization as per SEBI Circular dated Oct 6, 2017"]]
        bucket = category_map.get(str(category_label).strip()) if category_label else None
        if bucket is None:
            continue
        result[isin] = (int(rank), bucket)

    return result


def fetch_amfi_categorization(
    xlsx_url: str, session: requests.Session | None = None, timeout: float = 30.0
) -> dict[str, tuple[int, CapBucket]]:
    """xlsx_url must be the current edition's direct .xlsx link, found on
    AMFI_CATEGORIZATION_PAGE_URL - there's no stable "latest" URL to fetch
    blindly (see that constant's comment)."""
    session = session or requests.Session()
    response = session.get(
        xlsx_url, headers={"User-Agent": DEFAULT_USER_AGENT}, timeout=timeout
    )
    response.raise_for_status()
    return parse_amfi_categorization(response.content)


def tag_cap_buckets(
    universe: list[UniverseEntry], amfi_by_isin: dict[str, tuple[int, CapBucket]]
) -> None:
    """Mutates each entry's cap_bucket/cap_rank in place from the AMFI
    mapping. Entries with no AMFI match (not in AMFI's published universe
    at all - typically genuinely tiny/illiquid names) are left untagged,
    not excluded here; the liquidity filter (SPEC.md section 5) is a
    separate, later step, not this one's job.
    """
    matched = 0
    for entry in universe:
        amfi_entry = amfi_by_isin.get(entry.isin)
        if amfi_entry is None:
            continue
        entry.cap_rank, entry.cap_bucket = amfi_entry
        matched += 1
    logger.info("%d of %d universe entries matched to an AMFI cap bucket", matched, len(universe))


# --- Basic exclusions (partial - see module docstring) ----------------------


def exclude_non_standard_instruments(universe: list[UniverseEntry]) -> None:
    """Heuristic-only REIT/InvIT/investment-trust exclusion by whole-word
    name pattern - see NON_STANDARD_NAME_PATTERN's comment for why this
    isn't authoritative and, per a real audit against the full live
    universe (docs/decisions/0002-etf-reit-invit-heuristic-audit.md),
    doesn't need to be: real REITs/InvITs aren't in this data source at
    all. Mutates `excluded`/`exclusion_reason` in place; entries already
    excluded for another reason are left alone (first exclusion reason
    wins).
    """
    for entry in universe:
        if entry.excluded:
            continue
        if NON_STANDARD_NAME_PATTERN.search(entry.name):
            entry.excluded = True
            entry.exclusion_reason = (
                "name matches Trust/REIT/InvIT heuristic (unverified) - likely an "
                "investment/holding company rather than a literal REIT/InvIT, see decision 0002"
            )


def exclude_likely_financials_and_utilities(universe: list[UniverseEntry]) -> None:
    """Name-heuristic pre-filter for SPEC.md section 5's Financial
    Services and Utilities exclusions - see FINANCIAL_NAME_MARKERS'
    comment for why this exists and what it isn't (not authoritative;
    a real sector-based exclusion still belongs downstream once each
    company's actual sector is known from a fetch that succeeds).
    """
    for entry in universe:
        if entry.excluded:
            continue
        upper_name = f" {entry.name.upper()} "
        if any(marker in upper_name for marker in FINANCIAL_NAME_MARKERS):
            entry.excluded = True
            entry.exclusion_reason = "name matches Financial Services heuristic (unverified)"
        elif any(marker in upper_name for marker in UTILITY_NAME_MARKERS):
            entry.excluded = True
            entry.exclusion_reason = "name matches Utilities heuristic (unverified)"


def exclude_insufficient_listing_history(
    universe: list[UniverseEntry], as_of: date, min_days: int = 365
) -> None:
    """Excludes companies without at least one full fiscal year of listed
    history (SPEC.md section 1). Only applies to entries with a known NSE
    listing date - BSE-only entries (no listing date available from the
    feed this module uses) are left untagged by this specific rule rather
    than guessed at.
    """
    for entry in universe:
        if entry.excluded or entry.date_of_listing is None:
            continue
        if (as_of - entry.date_of_listing).days < min_days:
            entry.excluded = True
            entry.exclusion_reason = (
                f"listed {(as_of - entry.date_of_listing).days} days ago, "
                f"below the {min_days}-day minimum"
            )


def build_universe(amfi_xlsx_url: str, as_of: date) -> list[UniverseEntry]:
    """Convenience one-call pipeline: fetch NSE+BSE+AMFI, merge, tag cap
    buckets, apply all three exclusion passes in this module (non-standard
    instruments, financials/utilities, insufficient listing history).

    Lives here rather than duplicated in (or cross-imported between)
    scripts/pilot_fundamentals_fetch.py and
    scripts/full_universe_fundamentals_fetch.py - a script executed
    directly (`python scripts/foo.py`) only gets its own directory on
    sys.path, not the repo root, so `from scripts.other_script import ...`
    breaks in exactly that invocation style. Importing from the installed
    magicformula package doesn't have that problem.
    """
    logger.info("Fetching NSE equity list...")
    nse = fetch_nse_equity_list()
    logger.info("NSE: %d listings", len(nse))

    logger.info("Fetching BSE active equity list...")
    bse = fetch_bse_active_equity_list()
    logger.info("BSE: %d listings", len(bse))

    logger.info("Fetching AMFI categorization...")
    amfi = fetch_amfi_categorization(amfi_xlsx_url)
    logger.info("AMFI: %d categorized entities", len(amfi))

    entries = merge_universe(nse, bse)
    logger.info("Merged universe: %d entities", len(entries))

    tag_cap_buckets(entries, amfi)
    exclude_non_standard_instruments(entries)
    exclude_likely_financials_and_utilities(entries)
    exclude_insufficient_listing_history(entries, as_of=as_of)

    return entries
